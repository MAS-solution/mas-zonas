# Fuerza de venta nueva para los telefónicos del camión 700: N vendedores, cada cliente se visita el día que
# ya tiene en el maestro de Gescom (los que no tienen día toman el de su sector). Cada día se parte en N
# pedazos con cortes rectos norte-sur / este-oeste, los pedazos de los 5 días se emparejan por vendedor
# (totales parejos y siempre para el mismo lado de la ciudad) y cada vendedor-día lleva su orden de visita.
# Se corre desde datos/: python ../scripts/recupero_700.py [N]
import json, csv, sys, collections, math
import numpy as np
from scipy.optimize import linear_sum_assignment

N = int(sys.argv[1]) if len(sys.argv) > 1 else 4
LEJOS_KM = 15     # más lejos que esto del centro quedan fuera de la ruta (sierras, Monte Cristo...)
DIAS = ['Lunes', 'Martes', 'Miércoles', 'Jueves', 'Viernes']
OUT = 'C:/Users/usuario/Documents/restructura-zonas-611/'
asig = json.load(open('asignacion.json'))['16']          # para el día del sector de los que no tienen día
cm = json.load(open('cliente_manzana.json'))
todos = [r for r in csv.DictReader(open('clientes_esc.csv', encoding='utf-8')) if '700' in r['estado']]
sin_ubic = [r for r in todos if r['en_mapa'] != 'si']
cli = [r for r in todos if r['en_mapa'] == 'si']
DG = {}
for r in json.load(open('dias_gescom.json', encoding='utf-8')):
    for d, k in enumerate(('l', 'm', 'x', 'j', 'vi')):
        if r[k]: DG[r['codigo']] = d
for r in cli:
    r['lat'] = float(r['lat']); r['lng'] = float(r['lng'])
    r['x'] = r['lng'] * 95.0; r['y'] = r['lat'] * 110.5
    r['dia_origen'] = 'maestro' if r['codigo'] in DG else 'sector'
    r['dia'] = DG[r['codigo']] if r['codigo'] in DG else asig[cm[r['codigo']]][1] - 1
med = np.median([[r['x'], r['y']] for r in cli], axis=0)
for r in cli: r['km_centro'] = math.hypot(r['x'] - med[0], r['y'] - med[1])
lejos = sorted([r for r in cli if r['km_centro'] > LEJOS_KM], key=lambda r: -r['km_centro'])
cli = [r for r in cli if r['km_centro'] <= LEJOS_KM]
P = np.array([[r['x'], r['y']] for r in cli])
CENTRO = P.mean(0)

def cortes(idx, k):
    """Parte los clientes idx en k grupos con cortes rectos en el hueco más ancho (margen 3% de clientes)."""
    if k == 1: return [idx]
    Q = P[idx]; eje = int(np.argmax(Q.std(0)))
    o = np.argsort(Q[:, eje]); xs = Q[o, eje]; k1 = k // 2; meta = len(idx) * k1 / k
    tol = max(2, len(idx) * 0.03)
    cand = [i for i in range(len(o) - 1) if abs(i + 1 - meta) <= tol and xs[i + 1] > xs[i]] or [int(round(meta)) - 1]
    i = max(cand, key=lambda i: (xs[i + 1] - xs[i], -abs(i + 1 - meta)))
    return cortes(idx[o[:i + 1]], k1) + cortes(idx[o[i + 1:]], k - k1)

piezas = []
for d in range(5):
    idx = np.array([i for i, r in enumerate(cli) if r['dia'] == d])
    fila = []
    for g in cortes(idx, N):
        c = P[g].mean(0) - CENTRO
        fila.append((g, math.atan2(c[1], c[0])))
    piezas.append(fila)
# emparejar: el día con más clientes fija el orden por ángulo, los demás se asignan buscando totales parejos
base = max(range(5), key=lambda d: sum(len(g) for g, _ in piezas[d]))
vend = [[] for _ in range(N)]; ang = np.zeros(N); tot = np.zeros(N)
for v, (g, a) in enumerate(sorted(piezas[base], key=lambda p: p[1])):
    vend[v].append((base, g)); ang[v] = a; tot[v] = len(g)
for d in sorted(set(range(5)) - {base}, key=lambda d: -np.ptp([len(g) for g, _ in piezas[d]])):
    sz = np.array([len(g) for g, _ in piezas[d]]); an = np.array([a for _, a in piezas[d]])
    da = np.abs((an[:, None] - ang[None, :] + np.pi) % (2 * np.pi) - np.pi)
    cost = (tot[None, :] + sz[:, None] - tot.mean() - sz.mean()) ** 2 + 30.0 * da
    for k, v in zip(*linear_sum_assignment(cost)): vend[v].append((d, piezas[d][k][0])); tot[v] += sz[k]

km = lambda a, b: math.hypot(a['x'] - b['x'], a['y'] - b['y'])
def ordenar(pts):
    """Arranca por el cliente más cercano al centro de la ciudad, va al más cercano y destraba cruces (2-opt)."""
    if len(pts) < 3: return pts[:]
    c = {'x': CENTRO[0], 'y': CENTRO[1]}
    o = [min(pts, key=lambda p: km(p, c))]; libre = [p for p in pts if p is not o[0]]
    while libre:
        u = o[-1]; j = min(range(len(libre)), key=lambda j: km(u, libre[j])); o.append(libre.pop(j))
    mejora = True
    while mejora:
        mejora = False
        for i in range(len(o) - 2):
            for j in range(i + 2, len(o)):
                a, b, cc = o[i], o[i + 1], o[j]; dd = o[j + 1] if j + 1 < len(o) else None
                antes = km(a, b) + (km(cc, dd) if dd else 0); desp = km(a, cc) + (km(b, dd) if dd else 0)
                if desp < antes - 1e-9: o[i + 1:j + 1] = o[i + 1:j + 1][::-1]; mejora = True
    return o

rutas = []   # (vendedor, dia, [clientes en orden], km)
for v in range(N):
    for d, g in sorted(vend[v], key=lambda t: t[0]):
        o = ordenar([cli[i] for i in g])
        rutas.append((v + 1, d, o, 1.3 * sum(km(o[i - 1], o[i]) for i in range(1, len(o)))))

# ---------- Excel ----------
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
wb = Workbook(); ws = wb.active; ws.title = 'Resumen'
ws.append(['Vendedor'] + DIAS + ['Semana', 'Km semana (aprox.)'])
cnt = collections.defaultdict(int); kms = collections.defaultdict(float)
for v, d, o, k in rutas: cnt[(v, d)] = len(o); kms[v] += k
for v in range(1, N + 1):
    ws.append([f'Vendedor {v}'] + [cnt[(v, d)] for d in range(5)] + [sum(cnt[(v, d)] for d in range(5)), round(kms[v], 1)])
ws.append(['Total'] + [sum(cnt[(v, d)] for v in range(1, N + 1)) for d in range(5)] + [len(cli), round(sum(kms.values()), 1)])
for c in ws[ws.max_row]: c.font = Font(bold=True)
ws.append([]); ws.append([f'{len(cli)} telefónicos del camión 700 con ubicación. Día = el del maestro de Gescom '
                          f'({sum(r["dia_origen"] == "maestro" for r in cli)}); sin día en el maestro, el de su sector '
                          f'({sum(r["dia_origen"] == "sector" for r in cli)}). Km en línea recta × 1,3, del primero al último.'])
ws.append([f'Más de {LEJOS_KM} km del centro (fuera de la ruta): {len(lejos)}, en la hoja "Fuera de ruta".'])
ws.append([f'Sin ubicación (no se pueden rutear): {len(sin_ubic)}, en la hoja "Sin ubicación".'])
H = ['Día', 'Orden', 'Código', 'Cliente', 'Vendedor actual', 'Segmento', 'Compras camión 700', 'Día según', 'Latitud', 'Longitud', 'Mapa']
for v in range(1, N + 1):
    w = wb.create_sheet(f'Vendedor {v}'); w.append(H)
    for vv, d, o, k in rutas:
        if vv != v: continue
        for i, r in enumerate(o, 1):
            w.append([DIAS[d], i, r['codigo'], r['cliente'], r['vendedor'], r['segmento'], int(r['comp_700'] or 0),
                      r['dia_origen'], r['lat'], r['lng'], f'https://www.google.com/maps?q={r["lat"]},{r["lng"]}'])
w = wb.create_sheet('Fuera de ruta'); w.append(['Código', 'Cliente', 'Vendedor actual', 'Segmento', 'Compras camión 700', 'Día', 'Km al centro', 'Mapa'])
for r in lejos: w.append([r['codigo'], r['cliente'], r['vendedor'], r['segmento'], int(r['comp_700'] or 0), DIAS[r['dia']], round(r['km_centro'], 1), f'https://www.google.com/maps?q={r["lat"]},{r["lng"]}'])
w = wb.create_sheet('Sin ubicación'); w.append(['Código', 'Cliente', 'Vendedor actual', 'Segmento', 'Compras camión 700'])
for r in sin_ubic: w.append([r['codigo'], r['cliente'], r['vendedor'], r['segmento'], int(r['comp_700'] or 0)])
for w in wb.worksheets:
    for c in w[1]: c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='14213D'); c.alignment = Alignment(wrap_text=True, vertical='center')
    w.freeze_panes = 'A2'
    for col in w.columns:
        w.column_dimensions[col[0].column_letter].width = max(9, min(42, max(len(str(c.value or '')) for c in col[:60]) + 2))
wb['Resumen'].column_dimensions['A'].width = 14
wb.save(OUT + f'ruta_recupero_700_{N}vend.xlsx')

# ---------- Mapa (HTML suelto, con datos de clientes: no publicar) ----------
datos = [{'v': v, 'd': d, 'km': round(k, 1), 'p': [[round(r['lat'], 6), round(r['lng'], 6), r['codigo'], r['cliente']] for r in o]} for v, d, o, k in rutas]
html = open(__file__.replace('recupero_700.py', 'recupero_700_mapa.html'), encoding='utf-8').read()
open(OUT + f'ruta_recupero_700_{N}vend.html', 'w', encoding='utf-8').write(
    html.replace('/*DATOS*/[]', json.dumps(datos, ensure_ascii=False)).replace('/*N*/4', str(N)))

for v in range(1, N + 1):
    print(f'V{v}', [cnt[(v, d)] for d in range(5)], sum(cnt[(v, d)] for d in range(5)), f'{kms[v]:.0f} km')
print('fuera de ruta', len(lejos), 'sin ubicación', len(sin_ubic))
